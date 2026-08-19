import unicodedata
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from datetime import datetime

COLUNA_TECNICO     = 1   
COLUNA_EMPRESA     = 2   
COLUNA_SERVICO     = 3   
COLUNA_BANCO       = 4   
COLUNA_DT_SOLICITA = 5   
COLUNA_DT_STATUS   = 6   
COLUNA_STATUS      = 7   
COLUNA_OBS         = 8   

LINHA_INICIO_DADOS = 3  

def _encontrar_aba(workbook, *palavras_chave):
    for nome in workbook.sheetnames:
        n = nome.lower()
        if all(p in n for p in palavras_chave):
            return workbook[nome]
    return None


def _encontrar_aba_pendente(workbook):
    for nome in workbook.sheetnames:
        if 'pendente' in nome.lower():
            return workbook[nome]
    return None


def _formatar_data(valor_celula):
    if isinstance(valor_celula, datetime):
        return valor_celula.strftime('%d/%m/%Y')
    return str(valor_celula or '').strip()


def ler_pendentes_do_excel(caminho_arquivo):
    try:
        wb = openpyxl.load_workbook(caminho_arquivo)
    except Exception as e:
        print(f'[planilha] Erro ao abrir arquivo: {e}')
        return []

    aba = _encontrar_aba_pendente(wb)
    if not aba:
        print('[planilha] Aba Pendente não encontrada. Abas disponíveis:', wb.sheetnames)
        return []

    registros = []
    for num_linha in range(LINHA_INICIO_DADOS, aba.max_row + 1):
        tecnico = aba.cell(num_linha, COLUNA_TECNICO).value
        empresa = aba.cell(num_linha, COLUNA_EMPRESA).value
        servico = aba.cell(num_linha, COLUNA_SERVICO).value
        banco   = aba.cell(num_linha, COLUNA_BANCO).value
        status  = aba.cell(num_linha, COLUNA_STATUS).value

        if not any([tecnico, empresa, servico, banco, status]):
            continue

        registros.append({
            'tecnico':     str(tecnico or '').strip(),
            'empresa':     str(empresa or '').strip(),
            'servico':     str(servico or '').strip(),
            'banco':       str(banco   or '').strip(),
            'dt_solicita': _formatar_data(aba.cell(num_linha, COLUNA_DT_SOLICITA).value),
            'dt_status':   _formatar_data(aba.cell(num_linha, COLUNA_DT_STATUS).value),
            'status':      str(status  or '').strip(),
            'obs':         str(aba.cell(num_linha, COLUNA_OBS).value or '').strip(),
        })

    return registros


def exportar_pendentes_para_excel(registros, caminho_destino, titulo_aba='Pendente'):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = titulo_aba[:31]  # limite do Excel pro nome da aba

    cabecalhos = ['TÉCNICO', 'EMPRESA', 'SERVIÇO', 'BANCO', 'DT.SOLICITA', 'DT.STATUS', 'STATUS', 'OBS']
    estilo_cab = Font(bold=True, color='FFFFFF')
    fundo_cab  = PatternFill(fill_type='solid', fgColor='0A1628')

    for col, texto in enumerate(cabecalhos, 1):
        celula = ws.cell(1, col, texto)
        celula.font      = estilo_cab
        celula.fill      = fundo_cab
        celula.alignment = Alignment(horizontal='center', vertical='center')

    for idx, reg in enumerate(registros, 2):
        ws.cell(idx, COLUNA_TECNICO).value     = reg.get('tecnico', '')
        ws.cell(idx, COLUNA_EMPRESA).value     = reg.get('empresa', '')
        ws.cell(idx, COLUNA_SERVICO).value     = reg.get('servico', '')
        ws.cell(idx, COLUNA_BANCO).value       = reg.get('banco', '')
        ws.cell(idx, COLUNA_DT_SOLICITA).value = reg.get('dt_solicita', '')
        ws.cell(idx, COLUNA_DT_STATUS).value   = reg.get('dt_status', '')
        ws.cell(idx, COLUNA_STATUS).value      = reg.get('status', '')
        ws.cell(idx, COLUNA_OBS).value         = reg.get('obs', '')

    larguras = [14, 30, 16, 20, 14, 14, 26, 40]
    for col, larg in enumerate(larguras, 1):
        ws.column_dimensions[ws.cell(1, col).column_letter].width = larg

    wb.save(caminho_destino)
    return caminho_destino

def ler_planilha_completa(caminho_arquivo):
    try:
        wb = openpyxl.load_workbook(caminho_arquivo)
    except Exception as e:
        print(f'[planilha] Erro ao abrir arquivo: {e}')
        return {}

    print('[planilha] Abas encontradas:', wb.sheetnames)

    resultado = {
        'pendente':   [],
        'cobranca':   [],
        'pagamentos': [],
        'extrato':    [],
        'cancelados': [],
    }

    aba_p = _encontrar_aba_pendente(wb)
    if aba_p:
        resultado['pendente'] = _ler_aba_pendente_sheet(aba_p)

    aba_c = _encontrar_aba(wb, 'cobran')
    if aba_c:
        resultado['cobranca'] = _ler_aba_homologados(aba_c)

    aba_pg = _encontrar_aba(wb, 'homolog', 'pag')
    if aba_pg:
        resultado['pagamentos'] = _ler_aba_homologados(aba_pg)

    aba_ex = _encontrar_aba(wb, 'extrat')
    if aba_ex:
        resultado['extrato'] = _ler_aba_homologados(aba_ex)

    aba_can = _encontrar_aba(wb, 'cancel')
    if aba_can:
        resultado['cancelados'] = _ler_cancelados_sheet(aba_can)

    return resultado

def _e_linha_separadora(empresa, servico, banco):
    return not empresa and not servico and not banco


def _ler_aba_pendente_sheet(aba):
    registros = []
    for num_linha in range(LINHA_INICIO_DADOS, aba.max_row + 1):
        tecnico = aba.cell(num_linha, COLUNA_TECNICO).value
        empresa = aba.cell(num_linha, COLUNA_EMPRESA).value
        servico = aba.cell(num_linha, COLUNA_SERVICO).value
        banco   = aba.cell(num_linha, COLUNA_BANCO).value
        status  = aba.cell(num_linha, COLUNA_STATUS).value
        if not any([tecnico, empresa, servico, banco, status]):
            continue
        if _e_linha_separadora(empresa, servico, banco):
            continue
        registros.append({
            'tecnico':     str(tecnico or '').strip(),
            'empresa':     str(empresa or '').strip(),
            'servico':     str(servico or '').strip(),
            'banco':       str(banco   or '').strip(),
            'dt_solicita': _formatar_data(aba.cell(num_linha, COLUNA_DT_SOLICITA).value),
            'dt_status':   _formatar_data(aba.cell(num_linha, COLUNA_DT_STATUS).value),
            'status':      str(status  or '').strip(),
            'obs':         str(aba.cell(num_linha, COLUNA_OBS).value or '').strip(),
        })
    return registros


def _ler_cancelados_sheet(aba):
    registros = []
    for num_linha in range(LINHA_INICIO_DADOS, aba.max_row + 1):
        row = {campo: aba.cell(num_linha, col).value for campo, col in _COL_CAN.items()}
        if not any(row[c] for c in ('tecnico', 'empresa', 'servico', 'banco', 'status')):
            continue
        if _e_linha_separadora(row['empresa'], row['servico'], row['banco']):
            continue
        registros.append({
            'tecnico':       str(row['tecnico']       or '').strip(),
            'empresa':       str(row['empresa']       or '').strip(),
            'servico':       str(row['servico']       or '').strip(),
            'banco':         str(row['banco']         or '').strip(),
            'carteira':      str(row['carteira']      or '').strip(),
            'conta':         str(row['conta']         or '').strip(),
            'dt_solicita':   _formatar_data(row['dt_solicita']),
            'dt_status':     _formatar_data(row['dt_status']),
            'status':        str(row['status']        or 'CANCELADO').strip(),
            'retorno':       str(row['retorno']       or '').strip(),
            'ordem_servico': str(row['ordem_servico'] or '').strip(),
            'obs':           str(row['obs']           or '').strip(),
        })
    return registros

_COL_HOM = {
    'tecnico':     1,  
    'empresa':     2,  
    'servico':     3,  
    'banco':       4,  
    'dt_solicita': 5,  
    'dt_status':   6,  
    'status':      7,  
    'obs':         8,  
}

_COL_CAN = {
    'tecnico':       1,
    'empresa':       2,
    'servico':       3,
    'banco':         4,
    'carteira':      5,
    'conta':         6,
    'dt_solicita':   7,
    'dt_status':     8,
    'status':        9,
    'retorno':      10,
    'ordem_servico':11,
    'obs':          12,
}


def _ler_aba_homologados(aba):
    registros = []
    for num_linha in range(LINHA_INICIO_DADOS, aba.max_row + 1):
        row = {campo: aba.cell(num_linha, col).value for campo, col in _COL_HOM.items()}
        if not any(row[c] for c in ('tecnico', 'empresa', 'servico', 'banco', 'status')):
            continue
        if _e_linha_separadora(row['empresa'], row['servico'], row['banco']):
            continue
        registros.append({
            'tecnico':     str(row['tecnico']     or '').strip(),
            'empresa':     str(row['empresa']     or '').strip(),
            'servico':     str(row['servico']     or '').strip(),
            'banco':       str(row['banco']       or '').strip(),
            'dt_solicita': _formatar_data(row['dt_solicita']),
            'dt_status':   _formatar_data(row['dt_status']),
            'status':      str(row['status']      or 'HOMOLOGADO').strip(),
            'obs':         str(row['obs']         or '').strip(),
        })
    return registros


def ler_homologados_cobranca_do_excel(caminho_arquivo):
    try:
        wb = openpyxl.load_workbook(caminho_arquivo)
    except Exception as e:
        print(f'[planilha] Erro ao abrir arquivo: {e}')
        return []
    aba = _encontrar_aba(wb, 'cobran')
    if not aba:
        print('[planilha] Aba Homologados Cobrança não encontrada. Abas:', wb.sheetnames)
        return []
    return _ler_aba_homologados(aba)


def ler_homologados_pagamentos_do_excel(caminho_arquivo):
    try:
        wb = openpyxl.load_workbook(caminho_arquivo)
    except Exception as e:
        print(f'[planilha] Erro ao abrir arquivo: {e}')
        return []
    aba = _encontrar_aba(wb, 'pagamento') or _encontrar_aba(wb, 'pagfor') or _encontrar_aba(wb, 'pag')
    if not aba:
        print('[planilha] Aba Homologados Pagamentos não encontrada. Abas:', wb.sheetnames)
        return []
    return _ler_aba_homologados(aba)


def ler_homologados_extrato_do_excel(caminho_arquivo):
    try:
        wb = openpyxl.load_workbook(caminho_arquivo)
    except Exception as e:
        print(f'[planilha] Erro ao abrir arquivo: {e}')
        return []
    aba = _encontrar_aba(wb, 'extrat')
    if not aba:
        print('[planilha] Aba Extrato Cartão não encontrada. Abas:', wb.sheetnames)
        return []
    return _ler_aba_homologados(aba)


def ler_cancelados_do_excel(caminho_arquivo):
    try:
        wb = openpyxl.load_workbook(caminho_arquivo)
    except Exception as e:
        print(f'[planilha] Erro ao abrir arquivo: {e}')
        return []
    aba = _encontrar_aba(wb, 'cancel')
    if not aba:
        print('[planilha] Aba Cancelados não encontrada. Abas:', wb.sheetnames)
        return []
    registros = []
    for num_linha in range(LINHA_INICIO_DADOS, aba.max_row + 1):
        row = {campo: aba.cell(num_linha, col).value for campo, col in _COL_CAN.items()}
        if not any(row[c] for c in ('tecnico', 'empresa', 'servico', 'banco', 'status')):
            continue
        registros.append({
            'tecnico':       str(row['tecnico']       or '').strip(),
            'empresa':       str(row['empresa']       or '').strip(),
            'servico':       str(row['servico']       or '').strip(),
            'banco':         str(row['banco']         or '').strip(),
            'carteira':      str(row['carteira']      or '').strip(),
            'conta':         str(row['conta']         or '').strip(),
            'dt_solicita':   _formatar_data(row['dt_solicita']),
            'dt_status':     _formatar_data(row['dt_status']),
            'status':        str(row['status']        or 'CANCELADO').strip(),
            'retorno':       str(row['retorno']       or '').strip(),
            'ordem_servico': str(row['ordem_servico'] or '').strip(),
            'obs':           str(row['obs']           or '').strip(),
        })
    return registros

def _normalizar_texto(txt):
    txt = str(txt or '').strip().lower()
    txt = unicodedata.normalize('NFKD', txt).encode('ascii', 'ignore').decode('ascii')
    return txt

_MAPA_CAMPOS_ESCRITURAL = {
    'status':                 ('status', 'situacao'),
    'data_solicitacao':       ('data solicitacao', 'data da solicitacao', 'dt solicitacao', 'dt. solicita', 'dt.solicita', 'carimbo de data'),
    'empresa':                ('empresa', 'nome da empresa'),
    'cnpj':                   ('cnpj',),
    'responsavel':            ('responsavel', 'contato'),
    'fone_responsavel':       ('fone. resp.', 'fone resp.', 'fone resp', 'telefone resp', 'telefone', 'fone'),
    'email_responsavel':      ('e-mail resp.', 'e-mail resp', 'email resp.', 'email resp', 'e-mail', 'email'),
    'banco':                  ('banco',),
    'conta_corrente':         ('conta corrente',),
    'agencia':                ('agencia',),
    'convenio':               ('convenio',),
    'servico':                ('servico',),
    'carteira':               ('cart.', 'carteira'),
    'desc_carteira':          ('desc. carteira', 'desc carteira', 'descricao carteira', 'tipo de carteira'),
    'variacao':               ('variacao',),
    'juros':                  ('juros',),
    'multa':                  ('multa',),
    'protesto':               ('protesto', 'instrucao'),
    'emissor':                ('emissor', 'emissao'),
    'ultima_numeracao':       ('ult. numeracao', 'ultima numeracao', 'ult numeracao'),
    'cnab':                   ('cnab',),
    'solicitante':            ('solicitante',),
    'nome_solicitante':       ('nome sol.', 'nome sol', 'nome solicitante'),
    'obs_emissao':            ('obs. emissao', 'obs emissao', 'observacao emissao', 'observacao adicional', 'observacoes', 'obs'),
    'solicitante_pagto':      ('solicitante pagto',),
    'nome_solicitante_pagto': ('nome sol. pagto', 'nome sol pagto', 'nome solicitante pagto'),
    'modalidades':            ('modalidades',),
    'obs_pagto':              ('obs. pagto', 'obs pagto', 'observacao pagto'),
    'tecnico':                ('tecnico',),
    'dt_status':              ('dt. status', 'dt.status', 'data status', 'data do status'),
    'sequencial_nosso_numero': ('sequencial do nosso numero', 'sequencial nosso numero'),
    'quantos_dias':           ('quantos dias',),
    'anexos':                 ('anexos',),
}

_PARES_CAMPOS_DUPLICADOS = (
    ('solicitante', 'solicitante_pagto', ('solicitando',)),
    ('nome_solicitante', 'nome_solicitante_pagto', ('nome',)),
    ('obs_emissao', 'obs_pagto', ('observ', 'obs')),
)

_MIN_CAMPOS_CABECALHO = 4  


def _valor_para_texto(valor):
    if isinstance(valor, datetime):
        if valor.time() == datetime.min.time():
            return valor.strftime('%d/%m/%Y')
        return valor.strftime('%d/%m/%Y %H:%M:%S')
    if isinstance(valor, float):
        if valor.is_integer():
            return str(int(valor))
        return str(valor)
    return str(valor or '').strip()


def _mapear_colunas_escritural(aba, num_linha_cabecalho):
    cabecalhos = {}
    for col in range(1, aba.max_column + 1):
        texto = _normalizar_texto(aba.cell(num_linha_cabecalho, col).value)
        if texto:
            cabecalhos[col] = texto

    mapa = {}
    usadas = set()

    for campo, candidatos in _MAPA_CAMPOS_ESCRITURAL.items():
        for col, texto in cabecalhos.items():
            if col in usadas:
                continue
            if texto in candidatos:
                mapa[campo] = col
                usadas.add(col)
                break

    for campo, candidatos in _MAPA_CAMPOS_ESCRITURAL.items():
        if campo in mapa:
            continue
        for col, texto in cabecalhos.items():
            if col in usadas:
                continue
            if any(c in texto for c in candidatos):
                mapa[campo] = col
                usadas.add(col)
                break

    if 'status' not in mapa and 1 not in usadas and 1 not in cabecalhos:
        mapa['status'] = 1
        usadas.add(1)

    col_modalidades = mapa.get('modalidades')
    if col_modalidades:
        for campo_antes, campo_depois, marcadores in _PARES_CAMPOS_DUPLICADOS:
            for col, texto in sorted(cabecalhos.items()):
                if col in usadas:
                    continue
                if not any(m in texto for m in marcadores):
                    continue
                destino = campo_antes if col < col_modalidades else campo_depois
                if destino in mapa:
                    continue
                mapa[destino] = col
                usadas.add(col)

    return mapa


def _linha_eh_cabecalho_fantasma(aba, num_linha, mapa):
    col_data = mapa.get('data_solicitacao')
    if not col_data:
        return False
    valor = aba.cell(num_linha, col_data).value
    if not isinstance(valor, str):
        return False
    texto = _normalizar_texto(valor)
    return any(c in texto for c in _MAPA_CAMPOS_ESCRITURAL['data_solicitacao'])


def _detectar_cabecalho(aba, linhas_candidatas=(1, 2, 3)):
    melhor = (None, {})
    for linha in linhas_candidatas:
        if linha > aba.max_row:
            continue
        mapa = _mapear_colunas_escritural(aba, linha)
        if len(mapa) >= _MIN_CAMPOS_CABECALHO:
            return linha, mapa
        if len(mapa) > len(melhor[1]):
            melhor = (linha, mapa)
    return melhor

def ler_escritural_do_excel(caminho_arquivo):
    try:
        wb = openpyxl.load_workbook(caminho_arquivo, data_only=True)
    except Exception as e:
        print(f'[planilha] Erro ao abrir arquivo: {e}')
        return []

    registros = []
    for nome_aba in wb.sheetnames:
        aba = wb[nome_aba]
        linha_cab, mapa = _detectar_cabecalho(aba)
        if not mapa or linha_cab is None:
            print(f'[planilha] Aba "{nome_aba}" ignorada: nenhum cabeçalho reconhecido.')
            continue

        linha_inicio_dados = linha_cab + 1
        if linha_inicio_dados <= aba.max_row:
            mapa_linha_seguinte = _mapear_colunas_escritural(aba, linha_inicio_dados)
            if len(mapa_linha_seguinte) >= _MIN_CAMPOS_CABECALHO:
                linha_inicio_dados += 1

        qtd_antes = len(registros)
        qtd_fantasmas = 0
        for num_linha in range(linha_inicio_dados, aba.max_row + 1):
            if _linha_eh_cabecalho_fantasma(aba, num_linha, mapa):
                qtd_fantasmas += 1
                continue
            valores = {
                campo: aba.cell(num_linha, col).value
                for campo, col in mapa.items()
            }
            if not any(v not in (None, '') for v in valores.values()):
                continue
            reg = {}
            for campo in _MAPA_CAMPOS_ESCRITURAL:
                reg[campo] = _valor_para_texto(valores.get(campo))
            registros.append(reg)
        print(f'[planilha] Aba "{nome_aba}": cabeçalho na linha {linha_cab}, '
              f'{len(mapa)} campos mapeados, {len(registros) - qtd_antes} registros lidos'
              f'{f", {qtd_fantasmas} linha(s) de cabeçalho fantasma ignorada(s)" if qtd_fantasmas else ""}.')

    return registros
